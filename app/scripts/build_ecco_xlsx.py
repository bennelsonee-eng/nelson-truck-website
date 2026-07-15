"""build_ecco_xlsx.py — Excel of ECCO in-stock items across Titan + Nelson.

Pulls from:
  app/data/mysql_dumps/ecco_parts_master.csv  (2,088 ECCO parts catalogued)
  app/data/mysql_dumps/ecco_tte_inv.csv       (Titan side, warehouse 10)
  app/data/mysql_dumps/ecco_nte_inv.csv       (Nelson side, warehouses 1 + 2)

Output:
  app/data/reports/ecco_in_stock_titan_nelson.xlsx
"""
from __future__ import annotations

import csv
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DATA = Path(__file__).resolve().parent.parent / "data" / "mysql_dumps"
OUT = Path(__file__).resolve().parent.parent / "data" / "reports" / "ecco_in_stock_titan_nelson.xlsx"
ARIAL = "Arial"

WAREHOUSE_NAMES = {10: "SPO (Titan)", 1: "Portland (NTE)", 2: "Kent (NTE)"}


def D(v):
    if v in (None, "", "0", "0.00", "0.0"):
        return None
    try:
        d = Decimal(v)
        return d if d.is_finite() else None
    except (InvalidOperation, TypeError, ValueError):
        return None


def load_master(path: Path) -> dict[str, dict]:
    out: dict[str, dict] = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as fp:
        for r in csv.DictReader(fp):
            ou = (r.get("ourparts_num") or "").strip()
            if not ou or ou == "#":
                continue
            out[ou] = {
                "parts_num": (r.get("parts_num") or "").strip(),
                "description": (r.get("description") or "").strip(),
                "extra_desc": (r.get("extra_desc") or "").strip(),
                "p1": D(r.get("P1")), "p2": D(r.get("P2")), "p3": D(r.get("P3")),
                "p4": D(r.get("P4")), "p5": D(r.get("P5")),
                "weight": D(r.get("weight")),
                "location": (r.get("location") or "").strip(),
                "status": (r.get("status") or "").strip(),
                "supplier": (r.get("supplier") or "").strip(),
            }
    return out


def load_inv_combined(*paths: Path) -> dict[str, dict]:
    """Combine TTE + NTE inv into a single per-(ourparts_num, warehouse) record."""
    inv: dict[tuple[str, int], dict] = {}
    for path in paths:
        if not path.exists():
            continue
        with path.open("r", encoding="utf-8-sig", newline="") as fp:
            for r in csv.DictReader(fp):
                ou = (r.get("ourparts_num") or "").strip()
                if not ou:
                    continue
                try:
                    oh = int(float(r.get("onhand") or 0))
                    av = int(float(r.get("available") or 0))
                    wh = int(r.get("warehouse") or 0)
                except ValueError:
                    continue
                if wh not in WAREHOUSE_NAMES:
                    continue
                gc = D(r.get("gl_cost"))
                key = (ou, wh)
                if key not in inv:
                    inv[key] = {"on_hand": 0, "available": 0, "gl_cost": None}
                inv[key]["on_hand"] += oh
                inv[key]["available"] += av
                if gc is not None and (inv[key]["gl_cost"] is None or gc > inv[key]["gl_cost"]):
                    inv[key]["gl_cost"] = gc

    # Pivot: group by ourparts_num
    by_part: dict[str, dict] = defaultdict(lambda: {
        "warehouses": {}, "total_on_hand": 0, "total_available": 0, "gl_cost": None,
    })
    for (ou, wh), v in inv.items():
        by_part[ou]["warehouses"][wh] = v
        by_part[ou]["total_on_hand"] += v["on_hand"]
        by_part[ou]["total_available"] += v["available"]
        if v["gl_cost"] is not None and (by_part[ou]["gl_cost"] is None or v["gl_cost"] > by_part[ou]["gl_cost"]):
            by_part[ou]["gl_cost"] = v["gl_cost"]
    return dict(by_part)


def main() -> None:
    master = load_master(DATA / "ecco_parts_master.csv")
    inv = load_inv_combined(DATA / "ecco_tte_inv.csv", DATA / "ecco_nte_inv.csv")

    rows = []
    for ou, i in inv.items():
        if i["total_on_hand"] <= 0:
            continue
        m = master.get(ou, {})
        rows.append({
            "ourparts_num": ou,
            "parts_num": m.get("parts_num", ""),
            "description": m.get("description", ""),
            "extra_desc": m.get("extra_desc", ""),
            "wh_10": i["warehouses"].get(10, {}).get("on_hand"),
            "wh_1":  i["warehouses"].get(1, {}).get("on_hand"),
            "wh_2":  i["warehouses"].get(2, {}).get("on_hand"),
            "total_on_hand": i["total_on_hand"],
            "total_available": i["total_available"],
            "gl_cost": float(i["gl_cost"]) if i["gl_cost"] else None,
            "p1": float(m["p1"]) if m.get("p1") else None,
            "p2": float(m["p2"]) if m.get("p2") else None,
            "p3": float(m["p3"]) if m.get("p3") else None,
            "p4": float(m["p4"]) if m.get("p4") else None,
            "p5": float(m["p5"]) if m.get("p5") else None,
            "weight": float(m["weight"]) if m.get("weight") else None,
            "location": m.get("location", ""),
            "status": m.get("status", ""),
            "supplier": m.get("supplier", ""),
        })
    rows.sort(key=lambda r: -(r["total_on_hand"] or 0))
    print(f"in-stock SKUs (TTE+NTE wh 10/1/2): {len(rows)}, total units: {sum(r['total_on_hand'] for r in rows)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "In-Stock Items"

    header_font = Font(name=ARIAL, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="1F4E78")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cols = [
        ("ourparts_num", "Our Part #", 18, "left"),
        ("parts_num", "Mfr Part #", 18, "left"),
        ("description", "Description", 38, "left_wrap"),
        ("extra_desc", "Extra Desc", 22, "left_wrap"),
        ("wh_10", "SPO (Titan)", 11, "right"),
        ("wh_1",  "Portland", 10, "right"),
        ("wh_2",  "Kent", 10, "right"),
        ("total_on_hand", "Total On Hand", 12, "right"),
        ("total_available", "Available", 11, "right"),
        ("gl_cost", "GL Cost", 11, "currency"),
        ("p1", "P1 List", 11, "currency"),
        ("p2", "P2 Retail", 11, "currency"),
        ("p3", "P3 Jobber", 11, "currency"),
        ("p4", "P4 Dealer", 11, "currency"),
        ("p5", "P5 Cost", 11, "currency"),
        ("weight", "Wt (lb)", 9, "right"),
        ("location", "Loc", 8, "center"),
        ("status", "Status", 8, "center"),
        ("ecco_url", "ECCO Search Link", 22, "url"),
    ]

    for j, (_, h, w, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32

    for i, r in enumerate(rows, start=2):
        for j, (key, _, _, fmt) in enumerate(cols, start=1):
            if key == "ecco_url":
                pn = r["parts_num"] or r["ourparts_num"]
                c = ws.cell(
                    row=i, column=j,
                    value=f'=HYPERLINK("https://www.eccoesg.com/us/en/products/ProductSearch?q={pn}", "Search eccoesg.com")',
                )
                c.font = Font(name=ARIAL, color="0563C1", underline="single")
                c.alignment = center
            else:
                v = r.get(key)
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=ARIAL)
                if fmt == "currency":
                    c.number_format = "$#,##0.00;($#,##0.00);-"
                    c.alignment = right
                elif fmt == "right":
                    c.alignment = right
                elif fmt == "center":
                    c.alignment = center
                elif fmt == "left_wrap":
                    c.alignment = left_wrap
                else:
                    c.alignment = left
            c.border = border

    total_row = len(rows) + 2
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(name=ARIAL, bold=True)
    ws.cell(row=total_row, column=1).alignment = right
    ws.cell(row=total_row, column=1).fill = PatternFill("solid", start_color="DDEBF7")
    for col_idx in (5, 6, 7, 8, 9):
        letter = get_column_letter(col_idx)
        c = ws.cell(row=total_row, column=col_idx,
                    value=f"=SUM({letter}2:{letter}{total_row - 1})")
        c.font = Font(name=ARIAL, bold=True)
        c.alignment = right
        c.fill = PatternFill("solid", start_color="DDEBF7")
        c.border = border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{total_row - 1}"

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "ECCO Stock Across Titan + Nelson - Summary"
    ws2["A1"].font = Font(name=ARIAL, bold=True, size=14)
    ws2.merge_cells("A1:D1")

    last = len(rows) + 1
    labels = [
        ("Generated", '=TEXT(TODAY(), "yyyy-mm-dd")'),
        ("Brand", "ECCO (ECCO Safety Group — parent of ECCO + Code 3)"),
        ("Manufacturer URL", "https://www.eccoesg.com"),
        ("Master catalog PDF", "app/data/catalogs/ecco_master_2025.pdf (62 MB, downloaded 2026-05-14)"),
        ("Data source", "TigerTech MySQL: tte_parts_master + tte_inv_days (wh 10) + nte_inv_days (wh 1, 2)"),
        ("Total in-stock SKUs", f"=COUNTA('In-Stock Items'!A2:A{last})"),
        ("Total units on hand", f"=SUM('In-Stock Items'!H2:H{last})"),
        ("Spokane units (TTE)", f"=SUM('In-Stock Items'!E2:E{last})"),
        ("Portland units (NTE)", f"=SUM('In-Stock Items'!F2:F{last})"),
        ("Kent units (NTE)", f"=SUM('In-Stock Items'!G2:G{last})"),
        ("Total inventory @ GL cost",
         f"=SUMPRODUCT('In-Stock Items'!H2:H{last}, 'In-Stock Items'!J2:J{last})"),
        ("Total retail value @ P2",
         f"=SUMPRODUCT('In-Stock Items'!H2:H{last}, 'In-Stock Items'!L2:L{last})"),
    ]
    for i, (label, val) in enumerate(labels, start=3):
        ws2.cell(row=i, column=1, value=label).font = Font(name=ARIAL, bold=True)
        ws2.cell(row=i, column=1).alignment = right
        c = ws2.cell(row=i, column=2, value=val)
        c.font = Font(name=ARIAL)
        c.alignment = left
        if "value" in label.lower() or "@ GL" in label or "@ P2" in label:
            c.number_format = "$#,##0.00"

    ws2["A16"] = "Top 15 Highest-Stock SKUs"
    ws2["A16"].font = Font(name=ARIAL, bold=True, size=12)
    ws2["A16"].fill = PatternFill("solid", start_color="DDEBF7")
    ws2.merge_cells("A16:E16")

    for j, h in enumerate(["Our Part #", "Description", "On Hand", "GL Cost", "Distribution"], start=1):
        c = ws2.cell(row=17, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    for i, r in enumerate(rows[:15], start=18):
        ws2.cell(row=i, column=1, value=r["ourparts_num"]).font = Font(name=ARIAL)
        ws2.cell(row=i, column=2, value=r["description"]).font = Font(name=ARIAL)
        c = ws2.cell(row=i, column=3, value=r["total_on_hand"])
        c.font = Font(name=ARIAL); c.alignment = right
        c = ws2.cell(row=i, column=4, value=r["gl_cost"])
        c.font = Font(name=ARIAL); c.number_format = "$#,##0.00;($#,##0.00);-"; c.alignment = right
        # Distribution column
        dist_parts = []
        if r["wh_10"]:
            dist_parts.append(f"SPO:{r['wh_10']}")
        if r["wh_1"]:
            dist_parts.append(f"PDX:{r['wh_1']}")
        if r["wh_2"]:
            dist_parts.append(f"KEN:{r['wh_2']}")
        c = ws2.cell(row=i, column=5, value=" / ".join(dist_parts))
        c.font = Font(name=ARIAL); c.alignment = center
        for j in range(1, 6):
            ws2.cell(row=i, column=j).border = border

    for col, width in (("A", 18), ("B", 44), ("C", 12), ("D", 12), ("E", 24)):
        ws2.column_dimensions[col].width = width

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()

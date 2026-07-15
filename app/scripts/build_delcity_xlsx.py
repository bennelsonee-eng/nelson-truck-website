"""build_delcity_xlsx.py — Quick Excel of DelCity in-stock items for Titan.

Inputs (pulled via PHP bridge to TigerTech MySQL):
    app/data/mysql_dumps/delcity_parts_master.csv   (215 DelCity parts)
    app/data/mysql_dumps/delcity_inv_days.csv       (264 inv rows, often dup'd)

Output:
    app/data/reports/delcity_in_stock_titan.xlsx
"""
from __future__ import annotations

import csv
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DATA = Path(__file__).resolve().parent.parent / "data" / "mysql_dumps"
OUT = Path(__file__).resolve().parent.parent / "data" / "reports" / "delcity_in_stock_titan.xlsx"
ARIAL = "Arial"


def load_url_map(path: Path) -> dict[str, str]:
    """ourparts_num → delcity.net product URL (built from sitemap crossref)."""
    out: dict[str, str] = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as fp:
        for r in csv.DictReader(fp):
            ou = (r.get("ourparts_num") or "").strip()
            url = (r.get("delcity_url") or "").strip()
            if ou and url:
                out[ou] = url
    return out


def D(v):
    if v in (None, "", "0", "0.00", "0.0"):
        return None
    try:
        d = Decimal(v)
        return d if d.is_finite() else None
    except (InvalidOperation, TypeError, ValueError):
        return None


def load_master(path: Path) -> dict[str, dict]:
    out: dict[str, dict] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as fp:
        for r in csv.DictReader(fp):
            ou = (r.get("ourparts_num") or "").strip()
            if not ou or ou == "#":
                continue
            out[ou] = {
                "parts_num": (r.get("parts_num") or "").strip(),
                "description": (r.get("description") or "").strip(),
                "extra_desc": (r.get("extra_desc") or "").strip(),
                "p1": D(r.get("P1")), "p2": D(r.get("P2")), "p3": D(r.get("P3")),
                "p4": D(r.get("P4")), "p5": D(r.get("P5")),
                "weight": D(r.get("weight")),
                "location": (r.get("location") or "").strip(),
                "status": (r.get("status") or "").strip(),
                "supplier": (r.get("supplier") or "").strip(),
            }
    return out


def load_inv(path: Path) -> dict[str, dict]:
    inv: dict[str, dict] = defaultdict(lambda: {
        "on_hand": 0, "available": 0, "gl_cost": None, "warehouses": set(),
    })
    with path.open("r", encoding="utf-8-sig", newline="") as fp:
        for r in csv.DictReader(fp):
            ou = (r.get("ourparts_num") or "").strip()
            if not ou:
                continue
            try:
                oh = int(float(r.get("onhand") or 0))
                av = int(float(r.get("available") or 0))
                wh = int(r.get("warehouse") or 0)
            except ValueError:
                continue
            gc = D(r.get("gl_cost"))
            inv[ou]["on_hand"] += oh
            inv[ou]["available"] += av
            if gc is not None and (inv[ou]["gl_cost"] is None or gc > inv[ou]["gl_cost"]):
                inv[ou]["gl_cost"] = gc
            inv[ou]["warehouses"].add(wh)
    return dict(inv)


def main() -> None:
    master = load_master(DATA / "delcity_parts_master.csv")
    inv = load_inv(DATA / "delcity_inv_days.csv")
    url_map = load_url_map(DATA / "delcity_url_map.csv")
    print(f"loaded {len(url_map)} sitemap-mapped delcity.net URLs")

    rows = []
    for ou, i in inv.items():
        if i["on_hand"] <= 0:
            continue
        m = master.get(ou, {})
        rows.append({
            "ourparts_num": ou,
            "parts_num": m.get("parts_num", ""),
            "description": m.get("description", ""),
            "extra_desc": m.get("extra_desc", ""),
            "on_hand": i["on_hand"],
            "available": i["available"],
            "gl_cost": float(i["gl_cost"]) if i["gl_cost"] else None,
            "warehouses": ",".join(str(w) for w in sorted(i["warehouses"])),
            "p1": float(m["p1"]) if m.get("p1") else None,
            "p2": float(m["p2"]) if m.get("p2") else None,
            "p3": float(m["p3"]) if m.get("p3") else None,
            "p4": float(m["p4"]) if m.get("p4") else None,
            "p5": float(m["p5"]) if m.get("p5") else None,
            "weight": float(m["weight"]) if m.get("weight") else None,
            "location": m.get("location", ""),
            "status": m.get("status", ""),
            "supplier": m.get("supplier", ""),
        })
    rows.sort(key=lambda r: -r["on_hand"])
    print(f"in-stock SKUs: {len(rows)}, total units: {sum(r['on_hand'] for r in rows)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "In-Stock Items"

    header_font = Font(name=ARIAL, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="1F4E78")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cols = [
        ("ourparts_num", "Our Part #", 18, "left"),
        ("parts_num", "Supplier Part #", 18, "left"),
        ("description", "Description", 40, "left_wrap"),
        ("extra_desc", "Extra Desc", 24, "left_wrap"),
        ("on_hand", "On Hand", 10, "right"),
        ("available", "Available", 10, "right"),
        ("warehouses", "Whs", 8, "center"),
        ("gl_cost", "GL Cost", 11, "currency"),
        ("p1", "P1 List", 11, "currency"),
        ("p2", "P2 Retail", 11, "currency"),
        ("p3", "P3 Jobber", 11, "currency"),
        ("p4", "P4 Dealer", 11, "currency"),
        ("p5", "P5 Cost", 11, "currency"),
        ("weight", "Weight (lb)", 10, "right"),
        ("location", "Loc", 8, "center"),
        ("status", "Status", 8, "center"),
        ("supplier", "Supplier", 10, "center"),
        ("delcity_url", "delcity.net Product Page", 22, "url"),
    ]

    for j, (_, header, width, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[1].height = 32

    for i, r in enumerate(rows, start=2):
        for j, (key, _, _, fmt) in enumerate(cols, start=1):
            if key == "delcity_url":
                deep = url_map.get(r["ourparts_num"])
                if deep:
                    c = ws.cell(
                        row=i, column=j,
                        value=f'=HYPERLINK("{deep}", "View on delcity.net")',
                    )
                    c.font = Font(name=ARIAL, color="0563C1", underline="single")
                else:
                    pn = r["parts_num"] or r["ourparts_num"]
                    c = ws.cell(
                        row=i, column=j,
                        value=f'=HYPERLINK("https://www.delcity.net/store/search?searchTerm={pn}", "Search delcity.net")',
                    )
                    c.font = Font(name=ARIAL, color="999999", underline="single", italic=True)
                c.alignment = center
            else:
                v = r.get(key)
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=ARIAL)
                if fmt == "currency":
                    c.number_format = "$#,##0.00;($#,##0.00);-"
                    c.alignment = right
                elif fmt == "right":
                    c.alignment = right
                elif fmt == "center":
                    c.alignment = center
                elif fmt == "left_wrap":
                    c.alignment = left_wrap
                else:
                    c.alignment = left
            c.border = border

    total_row = len(rows) + 2
    tlabel = ws.cell(row=total_row, column=1, value="TOTAL")
    tlabel.font = Font(name=ARIAL, bold=True)
    tlabel.alignment = right
    tlabel.fill = PatternFill("solid", start_color="DDEBF7")
    tlabel.border = border
    for col_idx in (5, 6):
        letter = get_column_letter(col_idx)
        c = ws.cell(row=total_row, column=col_idx,
                    value=f"=SUM({letter}2:{letter}{total_row - 1})")
        c.font = Font(name=ARIAL, bold=True)
        c.alignment = right
        c.fill = PatternFill("solid", start_color="DDEBF7")
        c.border = border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{total_row - 1}"

    # ---- Summary sheet ----
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "DelCity In-Stock Summary - Titan Truck"
    ws2["A1"].font = Font(name=ARIAL, bold=True, size=14)
    ws2.merge_cells("A1:D1")
    ws2.row_dimensions[1].height = 24

    last = len(rows) + 1
    labels = [
        ("Generated", '=TEXT(TODAY(), "yyyy-mm-dd")'),
        ("Data source", "TigerTech MySQL: tte_parts_master + tte_inv_days (wh 10/1/2)"),
        ("Total in-stock SKUs", f"=COUNTA('In-Stock Items'!A2:A{last})"),
        ("Total units on hand", f"=SUM('In-Stock Items'!E2:E{last})"),
        ("Total inventory value @ GL cost",
         f"=SUMPRODUCT('In-Stock Items'!E2:E{last}, 'In-Stock Items'!H2:H{last})"),
        ("Total retail value @ P2",
         f"=SUMPRODUCT('In-Stock Items'!E2:E{last}, 'In-Stock Items'!J2:J{last})"),
    ]
    for i, (label, val) in enumerate(labels, start=3):
        ws2.cell(row=i, column=1, value=label).font = Font(name=ARIAL, bold=True)
        ws2.cell(row=i, column=1).alignment = right
        c = ws2.cell(row=i, column=2, value=val)
        c.font = Font(name=ARIAL)
        c.alignment = left
        if "value" in label.lower():
            c.number_format = "$#,##0.00"

    ws2["A10"] = "Top 10 Highest-Stock SKUs"
    ws2["A10"].font = Font(name=ARIAL, bold=True, size=12)
    ws2["A10"].fill = PatternFill("solid", start_color="DDEBF7")
    ws2.merge_cells("A10:D10")

    for j, h in enumerate(["Our Part #", "Description", "On Hand", "GL Cost"], start=1):
        c = ws2.cell(row=11, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    for i, r in enumerate(rows[:10], start=12):
        ws2.cell(row=i, column=1, value=r["ourparts_num"]).font = Font(name=ARIAL)
        ws2.cell(row=i, column=2, value=r["description"]).font = Font(name=ARIAL)
        c = ws2.cell(row=i, column=3, value=r["on_hand"])
        c.font = Font(name=ARIAL); c.alignment = right
        c = ws2.cell(row=i, column=4, value=r["gl_cost"])
        c.font = Font(name=ARIAL); c.number_format = "$#,##0.00;($#,##0.00);-"; c.alignment = right
        for j in range(1, 5):
            ws2.cell(row=i, column=j).border = border

    for col, width in (("A", 36), ("B", 44), ("C", 12), ("D", 12)):
        ws2.column_dimensions[col].width = width

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
